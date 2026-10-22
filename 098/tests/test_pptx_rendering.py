"""Tests for PPTX native rendering using python-pptx."""
from pathlib import Path
import pptx
from parsers.markdown_parser import MarkdownParser
from narrative.narrative_planner import NarrativePlanner
from layout.layout_solver import LayoutSolver
from renderers.pptx_renderer import PPTXRenderer


def test_pptx_renderer_creates_valid_deck(tmp_path):
    text = """# 工业级演示文稿
## 模块介绍
- 自动化渲染
- 物理质检
<!-- note: 这是演讲备注，用于真实验证。 -->
"""
    source_ir = MarkdownParser().parse_text(text)
    deck = NarrativePlanner().plan(source_ir)
    scene = LayoutSolver().solve_deck(deck)

    output_pptx = tmp_path / "test_deck.pptx"
    renderer = PPTXRenderer()
    renderer.render(scene, output_pptx)

    assert output_pptx.exists()
    assert output_pptx.stat().st_size > 1000  # Non-empty file

    # Inspect rendered pptx with python-pptx
    prs = pptx.Presentation(str(output_pptx))
    assert len(prs.slides) == len(scene.slides)

    # Check slide 0 has shapes
    slide_0 = prs.slides[0]
    assert len(slide_0.shapes) >= 3

    # Check notes on slides
    notes_found = False
    for s in prs.slides:
        if s.has_notes_slide:
            if "演讲备注" in s.notes_slide.notes_text_frame.text:
                notes_found = True
                break
    assert notes_found, "Speaker notes should be preserved on notes slide"
