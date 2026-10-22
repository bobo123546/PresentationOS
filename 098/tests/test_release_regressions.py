import json
import pytest
from pptx import Presentation
from parsers.markdown_parser import MarkdownParser
from narrative.narrative_planner import NarrativePlanner
from layout.layout_solver import LayoutSolver
from renderers.pptx_renderer import PPTXRenderer
from core.orchestrator import PresentationOrchestrator
from qa.native_qa_contract import NativeQAContract
from tools.validate_ir import validate_ir

def test_long_prose_and_tables_survive_actual_pptx_readback(tmp_path):
    body='Source detail without a number. '*40+'FINAL FACT 987654321'
    table='| Key | Value |\n| --- | --- |\n'+'\n'.join(f'| Row-{i} | {i} |' for i in range(14))
    source=tmp_path/'source.md'
    source.write_text('# Title\n## Evidence\n'+body+'\n\n'+table+'\n\nAFTER FIRST TABLE\n\n'+table+'\n',encoding='utf-8')
    result=PresentationOrchestrator(tmp_path/'data').compile(source,run_native_qa=False)
    assert result.success
    assert result.native_status=='NOT_ATTEMPTED'
    assert not result.gate_passed
    report=json.loads((result.evidence_dir/'provenance_audit.json').read_text())
    assert report['preserved'] and not report['missing_source_chunks']
    text=''.join(shape.text for slide in Presentation(str(result.pptx_path)).slides for shape in slide.shapes if shape.has_text_frame)
    assert 'FINALFACT987654321' in ''.join(text.split())
    assert 'AFTER FIRST TABLE' in text
    assert text.count('Row-13')==2
    manifest=json.loads((result.evidence_dir/'evidence_manifest.json').read_text())
    validate_ir('evidence_ir',manifest)
    assert manifest['status']=='LIMITED_TRIAL'
    assert manifest['source_snapshot']['files']
    assert manifest['command'] and manifest['exit_code']==0

def test_no_office_booleans_cannot_certify_unverified_native_checks():
    result=NativeQAContract.evaluate_run(True,True,True,True,True,True,package_valid=True,schema_valid=True)
    assert not result.is_release_ready

def test_renderer_rejects_unknown_node_and_overwrite(tmp_path):
    scene=LayoutSolver().solve_deck(NarrativePlanner().plan(MarkdownParser().parse_text('## Title\nBody')))
    scene.slides[0].nodes[0].kind='image'
    with pytest.raises(ValueError,match='Unsupported'): PPTXRenderer().render(scene,tmp_path/'bad.pptx')
    output=tmp_path/'existing.pptx';output.write_bytes(b'user bytes')
    with pytest.raises(FileExistsError): PPTXRenderer().render(scene,output)
    assert output.read_bytes()==b'user bytes'
