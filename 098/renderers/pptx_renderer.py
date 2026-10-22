"""PPTX Renderer converting ResolvedScene into a presentation file via python-pptx."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Optional, Tuple

import pptx
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from core.resolved_scene import ResolvedNode, ResolvedScene, ResolvedSlide


class PPTXRenderer:
    """Renders a ResolvedScene to a valid Microsoft PowerPoint (.pptx) file."""

    # 16:9 standard dimensions
    SLIDE_WIDTH_IN = 13.333
    SLIDE_HEIGHT_IN = 7.5

    def __init__(self) -> None:
        pass

    def render(self, scene: ResolvedScene, output_path: Path) -> Path:
        """Renders scene to .pptx presentation at output_path."""
        if output_path.exists():
            raise FileExistsError(f"Refusing to overwrite {output_path}")
        prs = pptx.Presentation()
        prs.slide_width = Inches(self.SLIDE_WIDTH_IN)
        prs.slide_height = Inches(self.SLIDE_HEIGHT_IN)

        # Blank layout is index 6 in standard templates
        blank_slide_layout = prs.slide_layouts[6]

        canvas_w = scene.canvas_width
        canvas_h = scene.canvas_height

        for r_slide in scene.slides:
            slide = prs.slides.add_slide(blank_slide_layout)

            # Sort nodes by z_index
            sorted_nodes = sorted(r_slide.nodes, key=lambda n: n.z_index)

            for node in sorted_nodes:
                self._render_node(slide, node, canvas_w, canvas_h)

            # Write speaker notes if present
            if r_slide.speaker_notes:
                notes_slide = slide.notes_slide
                text_frame = notes_slide.notes_text_frame
                text_frame.text = r_slide.speaker_notes

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("xb") as stream:
            prs.save(stream)
        return output_path

    def _render_node(
        self,
        slide: Any,
        node: ResolvedNode,
        canvas_w: float,
        canvas_h: float,
    ) -> None:
        """Converts canvas coordinates to inches and creates PPTX shape/textbox."""
        b = node.bounds
        left_in = Inches(b.x / canvas_w * self.SLIDE_WIDTH_IN)
        top_in = Inches(b.y / canvas_h * self.SLIDE_HEIGHT_IN)
        width_in = Inches(b.width / canvas_w * self.SLIDE_WIDTH_IN)
        height_in = Inches(b.height / canvas_h * self.SLIDE_HEIGHT_IN)

        if node.node_type in ("container", "shape"):
            # Background or card shape
            shape_type = MSO_SHAPE.RECTANGLE
            shape = slide.shapes.add_shape(shape_type, left_in, top_in, width_in, height_in)

            # Fill color
            if node.fill_color:
                shape.fill.solid()
                shape.fill.fore_color.rgb = self._parse_hex_color(node.fill_color)
            else:
                shape.fill.background()

            # Border line
            if node.border_color:
                shape.line.color.rgb = self._parse_hex_color(node.border_color)
                shape.line.width = Pt(1.0)
            else:
                shape.line.fill.background()

        elif node.node_type == "text":
            tx_box = slide.shapes.add_textbox(left_in, top_in, width_in, height_in)
            if node.source_id:
                tx_box.name = f"pres-source:{node.source_id}:{node.node_id}"
            tf = tx_box.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.05)
            tf.margin_right = Inches(0.05)
            tf.margin_top = Inches(0.05)
            tf.margin_bottom = Inches(0.05)

            p = tf.paragraphs[0]
            p.text = node.text or ""

            # Convert canvas pixel font size to PPTX Pt
            # standard 1080p canvas: 1080px = 7.5 inches = 540 pt, scale factor ~0.5
            font_size_pt = max(8.0, node.font_size * 0.5)
            p.font.size = Pt(font_size_pt)
            p.font.name = node.font_family or "Microsoft YaHei"

            if node.font_color:
                p.font.color.rgb = self._parse_hex_color(node.font_color)

            # Title or Section text bolding
            if node.font_size >= 24:
                p.font.bold = True

        else:
            raise ValueError(f"Unsupported render element: {node.node_type}")

    def _parse_hex_color(self, hex_code: str) -> RGBColor:
        """Parses #RRGGBB into RGBColor object."""
        clean = hex_code.lstrip("#")
        if len(clean) == 3:
            clean = "".join([c * 2 for c in clean])
        if len(clean) != 6:
            return RGBColor(255, 255, 255)
        try:
            r = int(clean[0:2], 16)
            g = int(clean[2:4], 16)
            b = int(clean[4:6], 16)
            return RGBColor(r, g, b)
        except ValueError:
            return RGBColor(255, 255, 255)
