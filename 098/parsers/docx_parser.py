"""DOCX parser extracting text hierarchy, tables and facts into SourceIR."""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Dict, List

import docx
from .markdown_parser import MarkdownParser


class DocxParser:
    """Parses Microsoft Word (.docx) files into standardized SourceIR."""

    def __init__(self) -> None:
        self.md_parser = MarkdownParser()

    def parse_file(self, file_path: Path) -> Dict[str, Any]:
        """Reads a .docx file and produces a SourceIR dictionary."""
        file_bytes = file_path.read_bytes()
        file_hash = hashlib.sha256(file_bytes).hexdigest()
        file_id = f"src_docx_{file_hash[:12]}"

        doc = docx.Document(str(file_path))
        chunks: List[Dict[str, Any]] = []
        facts: List[Dict[str, Any]] = []

        chunk_idx = 0
        fact_idx = 0
        current_section = "Document"

        for p_idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name.lower() if para.style else ""
            if "heading" in style_name or "title" in style_name:
                current_section = text
                cid = f"chk_{chunk_idx:04d}"
                chunks.append(
                    {
                        "chunk_id": cid,
                        "section_title": current_section,
                        "content_type": "heading",
                        "raw_text": text,
                        "line_range": [p_idx + 1, p_idx + 1],
                        "extracted_facts": [],
                    }
                )
                chunk_idx += 1
            else:
                cid = f"chk_{chunk_idx:04d}"
                extracted_facts = self.md_parser._extract_facts(text, cid, fact_idx)
                fact_idx += len(extracted_facts)
                facts.extend(extracted_facts)

                chunks.append(
                    {
                        "chunk_id": cid,
                        "section_title": current_section,
                        "content_type": "paragraph",
                        "raw_text": text,
                        "line_range": [p_idx + 1, p_idx + 1],
                        "extracted_facts": [f["fact_id"] for f in extracted_facts],
                    }
                )
                chunk_idx += 1

        # Tables
        for t_idx, table in enumerate(doc.tables):
            table_lines: List[str] = []
            for row in table.rows:
                row_vals = [cell.text.strip() for cell in row.cells]
                table_lines.append(" | ".join(row_vals))

            table_text = "\n".join(table_lines)
            if table_text:
                cid = f"chk_{chunk_idx:04d}"
                extracted_facts = self.md_parser._extract_facts(table_text, cid, fact_idx)
                fact_idx += len(extracted_facts)
                facts.extend(extracted_facts)

                chunks.append(
                    {
                        "chunk_id": cid,
                        "section_title": current_section,
                        "content_type": "table",
                        "raw_text": table_text,
                        "line_range": [t_idx + 1, t_idx + 1],
                        "extracted_facts": [f["fact_id"] for f in extracted_facts],
                    }
                )
                chunk_idx += 1

        return {
            "source_id": file_id,
            "file_name": file_path.name,
            "sha256": file_hash,
            "total_chunks": len(chunks),
            "chunks": chunks,
            "facts": facts,
        }
